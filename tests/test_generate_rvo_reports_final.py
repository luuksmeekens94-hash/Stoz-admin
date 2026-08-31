import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "generate-rvo-reports-final.py"
ACTIVITIES = ["A1.1", "A1.2", "A2.1", "A2.2", "A2.3", "A3.1", "A3.2", "A4.1", "A4.2", "A5.1", "A5.2", "A6.1", "A6.2"]


def load_generator():
    spec = importlib.util.spec_from_file_location("final_reports", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def entry(user, work_package, activity, hours, therapist=None, date="2026-08-31"):
    return {
        "date": date,
        "hours": hours,
        "status": "APPROVED",
        "withinCutoff": True,
        "user": user,
        "therapist": therapist,
        "workPackage": work_package,
        "activity": activity,
        "description": "Controleerbare testrealisatie.",
    }


def snapshot_fixture():
    entries = [
        entry("Manager Alpha", "WP1", "A1.1", 100),
        entry("Manager Beta", "WP1", "A1.1", 50),
        entry("Manager Gamma", "WP1", "A1.1", 22),
        entry("External Advisor", "WP1", "A1.1", 217.5),
        entry("Therapist Team", "WP2", "A2.1", 68),
        entry("Manager Alpha", "WP2", "A2.2", 15),
        entry("Manager Alpha", "WP5", "A5.1", 8),
        entry("Manager Alpha", "WP6", "A6.1", 12),
        entry("Manager Alpha", "WP3", "A3.1", 20),
        entry("Therapist Team", "WP3", "A3.1", 28, therapist="Test Therapist"),
        entry("Website Supplier", "WP2", "A2.3", 56),
    ]
    invoices = [
        {"number": "66", "amountExVat": 5180, "vatAmount": 1088, "confirmedBudgetLineId": "external-project-manager", "vatTreatment": "INCLUDED_CONFIRMED", "hasEvidence": True},
        {"number": "67", "amountExVat": 3150, "vatAmount": 661.5, "confirmedBudgetLineId": "external-project-manager", "vatTreatment": "INCLUDED_CONFIRMED", "hasEvidence": True},
        {"number": "71", "amountExVat": 2800, "vatAmount": 588, "confirmedBudgetLineId": "external-project-manager", "vatTreatment": "INCLUDED_CONFIRMED", "hasEvidence": True},
    ]
    return {
        "asOf": "2026-08-31",
        "applicant": "Fysiotherapie Fy-fit",
        "projectName": "Hybride Begrip",
        "approvedSubsidy": 39410,
        "entries": entries,
        "invoices": invoices,
        "trainings": [{"hours": 2, "presentAttendees": [f"Deelnemer {index}" for index in range(15)]}],
        "counts": {"clients": 0, "surveyResponses": 0},
        "activities": [{"code": code, "name": f"Activiteit {code}"} for code in ACTIVITIES],
        "planningVersion": {"allocations": []},
    }


class FinalReportGeneratorTest(unittest.TestCase):
    def test_activity_status_uses_parsed_dates(self):
        module = load_generator()
        self.assertEqual(module.activity_status(8, "01-08-2026", "28-02-2027", "31-08-2026", "31-08-2026"), "In uitvoering")
        self.assertEqual(module.activity_status(0, "01-03-2027", "31-08-2027", "31-08-2026", None), "Nog niet gestart")
        self.assertEqual(module.activity_status(8, "01-08-2025", "28-02-2026", "31-08-2026", "31-08-2026"), "Uitgevoerd / doorlopend na planfase")

    def test_generates_cutoff_reports_without_counting_therapist_attendance_in_model_b(self):
        module = load_generator()
        with tempfile.TemporaryDirectory() as tmp:
            snapshot = Path(tmp) / "snapshot.json"
            snapshot.write_text(json.dumps(snapshot_fixture()), encoding="utf-8")
            config = Path(tmp) / "report-config.json"
            config.write_text(json.dumps({
                "internalCostUsers": ["Manager Alpha", "Manager Beta", "Manager Gamma", "Therapist Team"],
                "websiteBuilderUsers": ["Website Supplier"],
            }), encoding="utf-8")
            outputs = module.build_reports(snapshot, Path(tmp) / "output", config)
            workbook = load_workbook(outputs["xlsx"], data_only=False)
            sheet = workbook["Aanvrager-Penvoerder"]

            project_hours = sum((sheet[f"E{row}"].value or 0) for row in range(16, 25))
            implementation_hours = sum((sheet[f"E{row}"].value or 0) for row in range(52, 61))
            learner_cost_inputs = [sheet[f"C{row}"].value for row in range(102, 111)]
            instructor_hours = sum((sheet[f"E{row}"].value or 0) for row in range(116, 125))

            self.assertEqual(project_hours, 172)
            self.assertEqual(implementation_hours, 103)
            self.assertTrue(all(value in (None, 0, "") for value in learner_cost_inputs))
            self.assertEqual(instructor_hours, 20)
            self.assertEqual(sheet["F75"].value, 2337.5)
            self.assertEqual((sheet["D69"].value, sheet["E69"].value), (100, 56))
            self.assertEqual(sheet["F76"].value, 1176)
            self.assertEqual((sheet["D87"].value, sheet["E87"].value), (12, 49))
            self.assertEqual((sheet["D88"].value, sheet["E88"].value), (12, 22))
            self.assertEqual(sheet["F139"].value, 645)
            self.assertIn("btw-plichtigheid ‘Nee’", sheet["B163"].value)
            self.assertIn("zonder dubbeltelling", sheet["B163"].value)
            self.assertIn("68 subsidiabele fysiotherapeuturen", sheet["B168"].value)
            self.assertIn("niet zonder onderliggende individuele berekening", sheet["B168"].value)

            doc = Document(outputs["docx"])
            text = "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
            self.assertIn("596,5", text)
            self.assertIn("WP5", text)
            self.assertIn("WP6", text)
            self.assertIn("0 cliënt", text)
            self.assertIn("A5.1 · Activiteit A5.1", text)
            self.assertIn("In uitvoering", text)
            self.assertIn("Nog niet gestart", text)
            self.assertNotIn("11 augustus 2026", text)


if __name__ == "__main__":
    unittest.main()
